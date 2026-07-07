import io
from datetime import timedelta
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib import messages
from django.db.models import Sum, Count, Avg, Q
from django.utils import timezone
from django.http import HttpResponse
from django.contrib.auth.models import User

from rooms.models import Room, RoomCategory
from bookings.models import Booking, Payment
from reviews.models import Review, ContactMessage
from reviews.ai import sentiment_summary
from .forms import RoomForm
from .models import Notification


@staff_member_required
def dashboard_view(request):
    total_users = User.objects.filter(is_staff=False).count()
    total_rooms = Room.objects.count()
    total_bookings = Booking.objects.count()
    revenue = Payment.objects.filter(status='success').aggregate(Sum('amount'))['amount__sum'] or 0

    occupied_rooms = Booking.objects.filter(status='confirmed').values('room').distinct().count()
    occupancy_rate = round((occupied_rooms / total_rooms) * 100, 1) if total_rooms else 0

    last_30 = timezone.now() - timedelta(days=30)
    recent_bookings = Booking.objects.filter(created_at__gte=last_30)

    # revenue trend for last 7 days (for Chart.js)
    revenue_labels, revenue_values = [], []
    for i in range(6, -1, -1):
        day = timezone.now().date() - timedelta(days=i)
        day_total = Payment.objects.filter(status='success', paid_at__date=day).aggregate(Sum('amount'))['amount__sum'] or 0
        revenue_labels.append(day.strftime('%d %b'))
        revenue_values.append(float(day_total))

    status_counts = Booking.objects.values('status').annotate(count=Count('id'))
    sentiment_stats = sentiment_summary(Review.objects.all())

    context = {
        'total_users': total_users,
        'total_rooms': total_rooms,
        'total_bookings': total_bookings,
        'revenue': revenue,
        'occupancy_rate': occupancy_rate,
        'recent_bookings_count': recent_bookings.count(),
        'revenue_labels': revenue_labels,
        'revenue_values': revenue_values,
        'status_counts': list(status_counts),
        'sentiment_stats': sentiment_stats,
        'latest_bookings': Booking.objects.select_related('room', 'user').all()[:8],
        'unread_messages': ContactMessage.objects.filter(is_read=False).count(),
    }
    return render(request, 'adminpanel/dashboard.html', context)


# ---------------- Room Management ----------------

@staff_member_required
def room_manage_view(request):
    rooms = Room.objects.select_related('category').all()
    return render(request, 'adminpanel/room_list.html', {'rooms': rooms})


@staff_member_required
def room_add_view(request):
    if request.method == 'POST':
        form = RoomForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, 'Room added successfully!')
            return redirect('adminpanel:room_manage')
    else:
        form = RoomForm()
    return render(request, 'adminpanel/room_form.html', {'form': form, 'title': 'Add Room'})


@staff_member_required
def room_edit_view(request, pk):
    room = get_object_or_404(Room, pk=pk)
    if request.method == 'POST':
        form = RoomForm(request.POST, request.FILES, instance=room)
        if form.is_valid():
            form.save()
            messages.success(request, 'Room updated successfully!')
            return redirect('adminpanel:room_manage')
    else:
        form = RoomForm(instance=room)
    return render(request, 'adminpanel/room_form.html', {'form': form, 'title': 'Edit Room'})


@staff_member_required
def room_delete_view(request, pk):
    room = get_object_or_404(Room, pk=pk)
    room.delete()
    messages.success(request, 'Room deleted.')
    return redirect('adminpanel:room_manage')


# ---------------- Booking Management ----------------

@staff_member_required
def booking_manage_view(request):
    bookings = Booking.objects.select_related('room', 'user').all()
    status = request.GET.get('status', '')
    if status:
        bookings = bookings.filter(status=status)
    return render(request, 'adminpanel/booking_list.html', {'bookings': bookings, 'status': status})


@staff_member_required
def booking_update_status(request, pk, new_status):
    booking = get_object_or_404(Booking, pk=pk)
    if new_status in dict(Booking.STATUS_CHOICES):
        booking.status = new_status
        booking.save()
        messages.success(request, f'Booking #{booking.id} marked as {new_status}.')

        # Let the customer know their booking status changed
        if booking.user.email:
            from django.core.mail import send_mail
            subject_map = {
                'confirmed': f'Your StayEase booking is confirmed! — {booking.room.name}',
                'cancelled': f'Your StayEase booking was cancelled — {booking.room.name}',
                'completed': f'Thanks for staying with StayEase — {booking.room.name}',
            }
            body_map = {
                'confirmed': (
                    f'Good news! The hotel has confirmed your booking.\n\n'
                    f'Room: {booking.room.name}\nCheck-in: {booking.check_in}\nCheck-out: {booking.check_out}\n\n'
                    f'We look forward to hosting you!'
                ),
                'cancelled': (
                    f'Your booking for {booking.room.name} ({booking.check_in} to {booking.check_out}) '
                    f'has been cancelled by the hotel. If you have questions, please contact support.'
                ),
                'completed': (
                    f'We hope you enjoyed your stay at {booking.room.name}! '
                    f'Please consider leaving a review to help other guests.'
                ),
            }
            if new_status in subject_map:
                send_mail(
                    subject_map[new_status], body_map[new_status],
                    'noreply@stayease.com', [booking.user.email], fail_silently=True,
                )
    return redirect('adminpanel:booking_manage')


# ---------------- User Management ----------------

@staff_member_required
def user_manage_view(request):
    users = User.objects.filter(is_staff=False).select_related('profile')
    return render(request, 'adminpanel/user_list.html', {'users': users})


@staff_member_required
def user_delete_view(request, pk):
    user = get_object_or_404(User, pk=pk)
    user.delete()
    messages.success(request, 'User deleted.')
    return redirect('adminpanel:user_manage')


@staff_member_required
def user_toggle_active(request, pk):
    user = get_object_or_404(User, pk=pk)
    user.is_active = not user.is_active
    user.save()
    return redirect('adminpanel:user_manage')


# ---------------- Review Management ----------------

@staff_member_required
def review_manage_view(request):
    reviews = Review.objects.select_related('user', 'room').all()
    sentiment = request.GET.get('sentiment', '')
    if sentiment:
        reviews = reviews.filter(sentiment=sentiment)
    return render(request, 'adminpanel/review_list.html', {
        'reviews': reviews, 'sentiment': sentiment, 'stats': sentiment_summary(Review.objects.all())
    })


@staff_member_required
def review_delete_view(request, pk):
    review = get_object_or_404(Review, pk=pk)
    review.delete()
    messages.success(request, 'Review deleted.')
    return redirect('adminpanel:review_manage')


# ---------------- Reports ----------------

@staff_member_required
def reports_view(request):
    context = {
        'total_bookings': Booking.objects.count(),
        'total_revenue': Payment.objects.filter(status='success').aggregate(Sum('amount'))['amount__sum'] or 0,
        'total_customers': User.objects.filter(is_staff=False).count(),
        'total_rooms': Room.objects.count(),
        'category_breakdown': RoomCategory.objects.annotate(booking_count=Count('rooms__bookings')),
    }
    return render(request, 'adminpanel/reports.html', context)


@staff_member_required
def export_bookings_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Bookings Report'

    headers = ['Booking ID', 'Guest', 'Room', 'Check-in', 'Check-out', 'Nights', 'Guests', 'Status', 'Total Amount']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='0F3D5C', end_color='0F3D5C', fill_type='solid')

    for b in Booking.objects.select_related('room', 'user').all():
        ws.append([
            str(b.booking_id), b.user.get_full_name() or b.user.username, b.room.name,
            b.check_in.strftime('%Y-%m-%d'), b.check_out.strftime('%Y-%m-%d'),
            b.nights, b.guests, b.get_status_display(), float(b.total_amount),
        ])

    for col in ws.columns:
        max_len = max(len(str(c.value)) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = max_len

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="bookings_report.xlsx"'
    return response


@staff_member_required
def export_revenue_pdf(request):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [Paragraph("<b>StayEase - Revenue Report</b>", styles['Title']), Spacer(1, 16)]

    data = [['Booking ID', 'Room', 'Guest', 'Status', 'Amount (Rs.)']]
    total = 0
    for b in Booking.objects.select_related('room', 'user').all()[:200]:
        data.append([str(b.id), b.room.name, b.user.username, b.get_status_display(), str(b.total_amount)])
        if b.status == 'confirmed' or b.status == 'completed':
            total += float(b.total_amount)

    table = Table(data, colWidths=[60, 140, 100, 90, 90])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f3d5c')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 20))
    elements.append(Paragraph(f"<b>Total Confirmed Revenue: Rs. {total:,.2f}</b>", styles['Heading3']))

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="revenue_report.pdf"'
    return response


# ---------------- Notifications (bell icon in navbar) ----------------

@staff_member_required
def notification_click(request, pk):
    """Mark a single notification read, then redirect to its linked page."""
    notif = get_object_or_404(Notification, pk=pk)
    notif.is_read = True
    notif.save()
    return redirect(notif.link or 'adminpanel:dashboard')


@staff_member_required
def notifications_mark_all_read(request):
    Notification.objects.filter(is_read=False).update(is_read=True)
    return redirect(request.META.get('HTTP_REFERER', 'adminpanel:dashboard'))
